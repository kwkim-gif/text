# ============================================================
# PPTX 텍스트 추출기 - 오탈자 검토용 GUI 프로그램 (Windows 11)
# ============================================================
# pip install python-pptx tkinterdnd2
# ============================================================

import os
import sys
import queue
import threading
import subprocess
import tkinter as tk
from tkinter import filedialog, ttk, messagebox

try:
    from tkinterdnd2 import TkinterDnD, DND_FILES
    DND_AVAILABLE = True
except ImportError:
    DND_AVAILABLE = False

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


# ── 색상 팔레트 ──────────────────────────────────────────────
COLOR_BG        = "#F3F3F3"
COLOR_CARD      = "#FFFFFF"
COLOR_ACCENT    = "#0067C0"
COLOR_ACCENT_H  = "#005BA4"
COLOR_ACCENT_DIM= "#CCE0F5"
COLOR_BORDER    = "#E0E0E0"
COLOR_TEXT      = "#1A1A1A"
COLOR_SUBTEXT   = "#6B6B6B"
COLOR_SUCCESS   = "#107C41"
COLOR_ERROR     = "#C42B1C"
COLOR_WARN      = "#CA5010"
COLOR_DROPZONE  = "#EBF3FB"
COLOR_DROP_ACT  = "#D0E8F8"
COLOR_ROW_ODD   = "#FFFFFF"
COLOR_ROW_EVEN  = "#F7F9FC"
COLOR_ROW_DONE  = "#EAF5EE"
COLOR_ROW_ERR   = "#FDF0EE"
COLOR_ROW_SEL   = "#D6E8FB"

# 각 파일의 상태값
ST_WAITING  = "대기 중"
ST_RUNNING  = "추출 중..."
ST_DONE     = "완료"
ST_ERROR    = "오류"


# ── 텍스트 추출 로직 ─────────────────────────────────────────

def extract_text_from_shape(shape):
    lines = []
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        for row in shape.table.rows:
            row_texts = [c.text.strip() for c in row.cells if c.text.strip()]
            if row_texts:
                lines.append("\t".join(row_texts))
    elif shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(r.text for r in para.runs).strip()
            if text:
                lines.append(text)
    return lines


def extract_text_from_slide(slide):
    all_lines = []
    def process(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                process(shape.shapes)
            else:
                all_lines.extend(extract_text_from_shape(shape))
    process(slide.shapes)
    return all_lines


def extract_pptx_to_txt(pptx_path, progress_callback=None):
    base_name  = os.path.splitext(os.path.basename(pptx_path))[0]
    output_dir = os.path.dirname(pptx_path)
    output_path = os.path.join(output_dir, f"{base_name}_오탈자검토.txt")

    prs = Presentation(pptx_path)
    total = len(prs.slides)
    empty_count = 0
    output_lines = []

    for idx, slide in enumerate(prs.slides, start=1):
        if progress_callback:
            progress_callback(idx, total)
        output_lines.append("=========================================")
        output_lines.append(f"[Slide {idx}]")
        output_lines.append("=========================================")
        texts = extract_text_from_slide(slide)
        if texts:
            output_lines.extend(texts)
        else:
            output_lines.append(f"[Slide {idx}] 추출된 텍스트가 없습니다.")
            empty_count += 1
        output_lines.append("")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(output_lines))

    return output_path, total, empty_count


# ── 파일 항목 데이터 클래스 ──────────────────────────────────

class FileItem:
    def __init__(self, path):
        self.path    = path
        self.name    = os.path.basename(path)
        self.status  = ST_WAITING
        self.slides  = ""
        self.message = ""
        self.output  = ""


# ── 파일 목록 테이블 위젯 ─────────────────────────────────────

class FileListWidget(tk.Frame):
    COLS = [
        ("번호",   40,  "center"),
        ("파일명", 280, "w"),
        ("슬라이드", 70, "center"),
        ("상태",   90,  "center"),
        ("결과",   120, "w"),
    ]

    def __init__(self, parent, on_select=None, **kwargs):
        super().__init__(parent, bg=COLOR_CARD, **kwargs)
        self._on_select = on_select
        self._items: list[FileItem] = []
        self._selected_idx = -1
        self._build()

    def _build(self):
        # 헤더
        header = tk.Frame(self, bg=COLOR_ACCENT)
        header.pack(fill="x")
        for col_name, width, anchor in self.COLS:
            tk.Label(
                header, text=col_name,
                bg=COLOR_ACCENT, fg="white",
                font=("Segoe UI", 9, "bold"),
                width=width // 7, anchor=anchor,
                padx=6, pady=6,
            ).pack(side="left")

        # 스크롤 가능 목록
        scroll_frame = tk.Frame(self, bg=COLOR_CARD)
        scroll_frame.pack(fill="both", expand=True)

        self._canvas = tk.Canvas(scroll_frame, bg=COLOR_CARD,
                                 highlightthickness=0, bd=0)
        scrollbar = ttk.Scrollbar(scroll_frame, orient="vertical",
                                  command=self._canvas.yview)
        self._canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")
        self._canvas.pack(side="left", fill="both", expand=True)

        self._list_frame = tk.Frame(self._canvas, bg=COLOR_CARD)
        self._canvas_window = self._canvas.create_window(
            (0, 0), window=self._list_frame, anchor="nw"
        )
        self._list_frame.bind("<Configure>", self._on_frame_configure)
        self._canvas.bind("<Configure>", self._on_canvas_configure)

        # 마우스 휠 스크롤
        self._canvas.bind("<MouseWheel>",
                          lambda e: self._canvas.yview_scroll(-1*(e.delta//120), "units"))

        self._row_frames = []
        self._render_empty()

    def _on_frame_configure(self, e):
        self._canvas.configure(scrollregion=self._canvas.bbox("all"))

    def _on_canvas_configure(self, e):
        self._canvas.itemconfig(self._canvas_window, width=e.width)

    def _render_empty(self):
        for w in self._list_frame.winfo_children():
            w.destroy()
        self._row_frames.clear()
        lbl = tk.Label(
            self._list_frame,
            text="PPTX 파일을 드래그하거나 [파일 추가] 버튼을 눌러 추가하세요",
            bg=COLOR_CARD, fg=COLOR_SUBTEXT,
            font=("Segoe UI", 10),
            pady=30,
        )
        lbl.pack(fill="x")

    def refresh(self):
        for w in self._list_frame.winfo_children():
            w.destroy()
        self._row_frames.clear()

        if not self._items:
            self._render_empty()
            return

        for i, item in enumerate(self._items):
            bg = COLOR_ROW_ODD if i % 2 == 0 else COLOR_ROW_EVEN
            if item.status == ST_DONE:
                bg = COLOR_ROW_DONE
            elif item.status == ST_ERROR:
                bg = COLOR_ROW_ERR
            if i == self._selected_idx:
                bg = COLOR_ROW_SEL

            row = tk.Frame(self._list_frame, bg=bg, cursor="hand2")
            row.pack(fill="x")
            self._row_frames.append(row)

            vals = [
                (str(i + 1),       40,  "center"),
                (item.name,        280, "w"),
                (item.slides,      70,  "center"),
                (item.status,      90,  "center"),
                (item.message,     120, "w"),
            ]
            # 상태 색상
            st_color = {
                ST_WAITING: COLOR_SUBTEXT,
                ST_RUNNING: COLOR_ACCENT,
                ST_DONE:    COLOR_SUCCESS,
                ST_ERROR:   COLOR_ERROR,
            }.get(item.status, COLOR_TEXT)

            for j, (val, width, anchor) in enumerate(vals):
                color = st_color if j == 3 else COLOR_TEXT
                tk.Label(
                    row, text=val,
                    bg=bg, fg=color,
                    font=("Segoe UI", 9),
                    width=width // 7, anchor=anchor,
                    padx=6, pady=7,
                ).pack(side="left")

            idx = i  # 클로저 캡처
            row.bind("<Button-1>", lambda e, n=idx: self._on_row_click(n))
            for child in row.winfo_children():
                child.bind("<Button-1>", lambda e, n=idx: self._on_row_click(n))

        # 구분선
        tk.Frame(self._list_frame, bg=COLOR_BORDER, height=1).pack(fill="x")

    def _on_row_click(self, idx):
        self._selected_idx = idx
        self.refresh()
        if self._on_select:
            self._on_select(self._items[idx])

    def add_paths(self, paths: list[str]):
        existing = {item.path for item in self._items}
        added = 0
        for p in paths:
            p = p.strip().strip("{}")  # tkinterdnd2 중괄호 제거
            if p.lower().endswith(".pptx") and p not in existing:
                self._items.append(FileItem(p))
                existing.add(p)
                added += 1
        if added:
            self.refresh()
        return added

    def remove_selected(self):
        if 0 <= self._selected_idx < len(self._items):
            self._items.pop(self._selected_idx)
            self._selected_idx = min(self._selected_idx, len(self._items) - 1)
            self.refresh()

    def clear_all(self):
        self._items.clear()
        self._selected_idx = -1
        self.refresh()

    def update_item(self, idx, **kwargs):
        if 0 <= idx < len(self._items):
            item = self._items[idx]
            for k, v in kwargs.items():
                setattr(item, k, v)
            self.refresh()

    @property
    def items(self):
        return self._items

    @property
    def selected(self):
        if 0 <= self._selected_idx < len(self._items):
            return self._items[self._selected_idx]
        return None


# ── 메인 앱 ──────────────────────────────────────────────────

class PPTXExtractorApp(TkinterDnD.Tk if DND_AVAILABLE else tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("PPTX 텍스트 추출기")
        self.geometry("780x640")
        self.minsize(640, 520)
        self.configure(bg=COLOR_BG)
        self.resizable(True, True)

        try:
            self.iconbitmap(default="")
        except Exception:
            pass

        self._last_output_folder = ""
        self._is_running = False
        self._queue = queue.Queue()   # 스레드 → 메인 스레드 통신용

        self._build_ui()
        self._center_window()

    def _center_window(self):
        self.update_idletasks()
        w, h = self.winfo_width(), self.winfo_height()
        x = (self.winfo_screenwidth()  - w) // 2
        y = (self.winfo_screenheight() - h) // 2
        self.geometry(f"{w}x{h}+{x}+{y}")

    # ── UI 구성 ───────────────────────────────────────────────

    def _build_ui(self):
        # 헤더
        header = tk.Frame(self, bg=COLOR_ACCENT, height=56)
        header.pack(fill="x")
        header.pack_propagate(False)
        tk.Label(header, text="📄  PPTX 텍스트 추출기",
                 bg=COLOR_ACCENT, fg="white",
                 font=("Segoe UI", 14, "bold"), padx=20
                 ).pack(side="left", fill="y")
        tk.Label(header, text="오탈자 검토용 TXT 일괄 변환",
                 bg=COLOR_ACCENT, fg="#C8DFEF",
                 font=("Segoe UI", 10), padx=4
                 ).pack(side="left", fill="y")

        body = tk.Frame(self, bg=COLOR_BG)
        body.pack(fill="both", expand=True, padx=20, pady=16)

        self._build_drop_zone(body)
        self._build_file_list(body)
        self._build_progress_area(body)
        self._build_bottom_bar()

    def _build_drop_zone(self, parent):
        """드래그 앤 드롭 안내 영역"""
        outer = tk.Frame(parent, bg=COLOR_BORDER)
        outer.pack(fill="x", pady=(0, 10))
        inner = tk.Frame(outer, bg=COLOR_DROPZONE, padx=0, pady=0)
        inner.pack(fill="x", padx=1, pady=1)

        self._drop_zone = tk.Frame(inner, bg=COLOR_DROPZONE, pady=18)
        self._drop_zone.pack(fill="x")

        tk.Label(self._drop_zone,
                 text="여기에 PPTX 파일을 드래그 앤 드롭하세요",
                 bg=COLOR_DROPZONE, fg=COLOR_ACCENT,
                 font=("Segoe UI", 11, "bold")).pack()
        tk.Label(self._drop_zone,
                 text="여러 파일을 한 번에 드롭하거나, 아래 [파일 추가] 버튼으로 선택할 수 있습니다",
                 bg=COLOR_DROPZONE, fg=COLOR_SUBTEXT,
                 font=("Segoe UI", 9)).pack(pady=(4, 0))

        if not DND_AVAILABLE:
            tk.Label(self._drop_zone,
                     text="⚠  tkinterdnd2 미설치 — 드래그 앤 드롭 비활성",
                     bg=COLOR_DROPZONE, fg=COLOR_WARN,
                     font=("Segoe UI", 8)).pack(pady=(6, 0))

        # 드래그 앤 드롭 이벤트 바인딩
        if DND_AVAILABLE:
            for widget in [self._drop_zone, outer, inner] + self._drop_zone.winfo_children():
                try:
                    widget.drop_target_register(DND_FILES)
                    widget.dnd_bind("<<Drop>>", self._on_drop)
                    widget.dnd_bind("<<DragEnter>>", self._on_drag_enter)
                    widget.dnd_bind("<<DragLeave>>", self._on_drag_leave)
                except Exception:
                    pass

        self._drop_outer = outer
        self._drop_inner = inner

    def _build_file_list(self, parent):
        """파일 목록 테이블 + 목록 조작 버튼"""
        outer = tk.Frame(parent, bg=COLOR_BORDER)
        outer.pack(fill="both", expand=True, pady=(0, 10))
        inner = tk.Frame(outer, bg=COLOR_CARD)
        inner.pack(fill="both", expand=True, padx=1, pady=1)

        # 목록 상단 툴바
        toolbar = tk.Frame(inner, bg=COLOR_CARD, pady=8, padx=10)
        toolbar.pack(fill="x")
        tk.Label(toolbar, text="파일 목록",
                 bg=COLOR_CARD, fg=COLOR_SUBTEXT,
                 font=("Segoe UI", 9, "bold")).pack(side="left")

        self._lbl_count = tk.Label(toolbar, text="0개",
                                   bg=COLOR_CARD, fg=COLOR_ACCENT,
                                   font=("Segoe UI", 9, "bold"))
        self._lbl_count.pack(side="left", padx=(6, 0))

        # 목록 조작 버튼 (우측)
        self._btn_clear = self._ghost_button(toolbar, "전체 삭제", self._on_clear_all)
        self._btn_clear.pack(side="right", padx=(6, 0))

        self._btn_remove = self._ghost_button(toolbar, "선택 삭제", self._on_remove)
        self._btn_remove.pack(side="right", padx=(6, 0))

        self._btn_add = self._accent_button(toolbar, "+ 파일 추가", self._on_add_files)
        self._btn_add.pack(side="right", padx=(6, 0))

        tk.Frame(inner, bg=COLOR_BORDER, height=1).pack(fill="x")

        # 파일 목록 테이블
        self._file_list = FileListWidget(
            inner, on_select=self._on_file_select, height=180
        )
        self._file_list.pack(fill="both", expand=True)

        # 드롭 이벤트를 파일 목록에도 연결
        if DND_AVAILABLE:
            try:
                self._file_list.drop_target_register(DND_FILES)
                self._file_list.dnd_bind("<<Drop>>", self._on_drop)
                self._file_list.dnd_bind("<<DragEnter>>", self._on_drag_enter)
                self._file_list.dnd_bind("<<DragLeave>>", self._on_drag_leave)
            except Exception:
                pass

    def _build_progress_area(self, parent):
        outer = tk.Frame(parent, bg=COLOR_BORDER)
        outer.pack(fill="x", pady=(0, 10))
        card = tk.Frame(outer, bg=COLOR_CARD, padx=14, pady=12)
        card.pack(fill="x", padx=1, pady=1)

        top = tk.Frame(card, bg=COLOR_CARD)
        top.pack(fill="x")
        self._lbl_status = tk.Label(top, text="파일을 추가하고 추출을 시작하세요.",
                                    bg=COLOR_CARD, fg=COLOR_SUBTEXT,
                                    font=("Segoe UI", 10))
        self._lbl_status.pack(side="left")

        self._lbl_progress_text = tk.Label(top, text="",
                                           bg=COLOR_CARD, fg=COLOR_SUBTEXT,
                                           font=("Segoe UI", 9))
        self._lbl_progress_text.pack(side="right")

        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("Accent.Horizontal.TProgressbar",
                        troughcolor=COLOR_BORDER, background=COLOR_ACCENT,
                        bordercolor=COLOR_BORDER, lightcolor=COLOR_ACCENT,
                        darkcolor=COLOR_ACCENT)
        self._progressbar = ttk.Progressbar(
            card, style="Accent.Horizontal.TProgressbar",
            orient="horizontal", mode="determinate")
        self._progressbar.pack(fill="x", ipady=3, pady=(8, 0))

    def _build_bottom_bar(self):
        bar = tk.Frame(self, bg=COLOR_BG)
        bar.pack(fill="x", padx=20, pady=(0, 16))

        self._btn_open_folder = tk.Button(
            bar, text="📂  저장 폴더 열기",
            bg=COLOR_CARD, fg=COLOR_TEXT,
            font=("Segoe UI", 10), relief="flat", bd=0,
            padx=14, pady=8, cursor="hand2", state="disabled",
            activebackground=COLOR_BORDER,
            command=self._on_open_folder,
        )
        self._btn_open_folder.pack(side="left")

        self._btn_run = self._accent_button(
            bar, "  텍스트 추출 시작  ", self._on_run_all, big=True)
        self._btn_run.pack(side="right")
        self._btn_run.configure(state="disabled")

    # ── 드래그 앤 드롭 이벤트 ─────────────────────────────────

    def _on_drag_enter(self, event):
        self._drop_zone.configure(bg=COLOR_DROP_ACT)
        self._drop_outer.configure(bg=COLOR_ACCENT)

    def _on_drag_leave(self, event):
        self._drop_zone.configure(bg=COLOR_DROPZONE)
        self._drop_outer.configure(bg=COLOR_BORDER)

    def _on_drop(self, event):
        self._drop_zone.configure(bg=COLOR_DROPZONE)
        self._drop_outer.configure(bg=COLOR_BORDER)

        # tkinterdnd2는 여러 파일을 공백 구분 또는 중괄호 묶음으로 전달
        raw = event.data
        paths = self._parse_drop_data(raw)
        added = self._file_list.add_paths(paths)
        self._update_count()
        if added:
            self._set_run_button_state()

    def _parse_drop_data(self, raw: str) -> list[str]:
        """tkinterdnd2 드롭 데이터 파싱 (경로에 공백 포함 가능)"""
        paths = []
        raw = raw.strip()
        i = 0
        while i < len(raw):
            if raw[i] == "{":
                end = raw.find("}", i)
                if end != -1:
                    paths.append(raw[i+1:end])
                    i = end + 1
                else:
                    break
            else:
                end = raw.find(" ", i)
                if end == -1:
                    paths.append(raw[i:])
                    break
                paths.append(raw[i:end])
                i = end + 1
        return [p.strip() for p in paths if p.strip()]

    # ── 버튼 이벤트 핸들러 ────────────────────────────────────

    def _on_add_files(self):
        paths = filedialog.askopenfilenames(
            title="PPTX 파일 선택 (다중 선택 가능)",
            filetypes=[("PowerPoint 파일", "*.pptx"), ("모든 파일", "*.*")],
        )
        if paths:
            self._file_list.add_paths(list(paths))
            self._update_count()
            self._set_run_button_state()

    def _on_remove(self):
        self._file_list.remove_selected()
        self._update_count()
        self._set_run_button_state()

    def _on_clear_all(self):
        self._file_list.clear_all()
        self._update_count()
        self._set_run_button_state()
        self._lbl_status.configure(text="파일을 추가하고 추출을 시작하세요.", fg=COLOR_SUBTEXT)
        self._progressbar["value"] = 0
        self._lbl_progress_text.configure(text="")
        self._btn_open_folder.configure(state="disabled")

    def _on_file_select(self, item: FileItem):
        pass  # 선택 시 추후 미리보기 확장 가능

    def _on_run_all(self):
        items = self._file_list.items
        if not items or self._is_running:
            return

        targets = [i for i, it in enumerate(items)
                   if it.status in (ST_WAITING, ST_ERROR)]
        if not targets:
            for item in items:
                item.status = ST_WAITING
                item.slides = ""
                item.message = ""
            targets = list(range(len(items)))
            self._file_list.refresh()

        self._is_running = True

        # Queue 초기화 (이전 잔여 메시지 제거)
        while not self._queue.empty():
            try:
                self._queue.get_nowait()
            except queue.Empty:
                break

        self._btn_run.configure(state="disabled", text="  추출 중...  ")
        self._btn_add.configure(state="disabled")
        self._btn_remove.configure(state="disabled")
        self._btn_clear.configure(state="disabled")
        self._btn_open_folder.configure(state="disabled")
        self._progressbar["value"] = 0
        self._last_output_folder = ""

        thread = threading.Thread(
            target=self._run_all_thread, args=(targets,), daemon=True)
        thread.start()

        # Queue 폴링 시작 (메인 스레드에서 50ms마다 메시지 처리)
        self.after(50, self._poll_queue)

    # ── Queue 폴링: 스레드 메시지를 메인 스레드에서 안전하게 처리 ──────────
    # 스레드는 절대 after()를 호출하지 않고, Queue에만 메시지를 넣는다.
    # 메인 스레드의 _poll_queue()가 꺼내서 UI를 업데이트한다.

    def _poll_queue(self):
        try:
            while True:
                msg = self._queue.get_nowait()
                kind = msg[0]

                if kind == "running":
                    _, idx, job_num, total_files, name = msg
                    self._file_list.update_item(idx, status=ST_RUNNING, message="")
                    self._lbl_status.configure(
                        text=f"[{job_num}/{total_files}]  {name}", fg=COLOR_ACCENT)

                elif kind == "progress":
                    _, pct, cur, total_slides, jn, tf = msg
                    self._progressbar["value"] = pct
                    self._lbl_progress_text.configure(
                        text=f"파일 {jn}/{tf}  |  Slide {cur}/{total_slides}  ({pct}%)")

                elif kind == "done":
                    _, idx, total_slides, out_path = msg
                    self._last_output_folder = os.path.dirname(out_path)
                    self._file_list.update_item(
                        idx, status=ST_DONE,
                        slides=str(total_slides),
                        message="저장 완료",
                        output=out_path)

                elif kind == "error":
                    _, idx, err_msg = msg
                    self._file_list.update_item(
                        idx, status=ST_ERROR, message=err_msg)

                elif kind == "all_done":
                    _, total_files = msg
                    self._on_all_done(total_files)
                    return  # 완료 → 폴링 중단

        except queue.Empty:
            pass

        # 아직 실행 중이면 계속 폴링
        if self._is_running:
            self.after(50, self._poll_queue)

    # ── 백그라운드 추출 스레드 (Queue에만 메시지 전송, UI 직접 접근 금지) ──

    def _run_all_thread(self, target_indices):
        items = self._file_list.items
        total_files = len(target_indices)

        try:
            for job_num, idx in enumerate(target_indices, start=1):
                item = items[idx]
                self._queue.put(("running", idx, job_num, total_files, item.name))

                def make_progress_cb(jn, tf):
                    def cb(cur, total_slides):
                        pct = int(((jn - 1) + cur / total_slides) / tf * 100)
                        self._queue.put(("progress", pct, cur, total_slides, jn, tf))
                    return cb

                try:
                    out, total_slides, empty = extract_pptx_to_txt(
                        item.path,
                        progress_callback=make_progress_cb(job_num, total_files),
                    )
                    self._queue.put(("done", idx, total_slides, out))
                except Exception as e:
                    self._queue.put(("error", idx, str(e)[:60]))

        except Exception as e:
            # 예기치 않은 스레드 전체 오류도 Queue로 전달
            self._queue.put(("error", 0, f"Thread error: {e}"))

        self._queue.put(("all_done", total_files))

    def _on_all_done(self, total_files):
        self._is_running = False
        items = self._file_list.items
        done  = sum(1 for it in items if it.status == ST_DONE)
        error = sum(1 for it in items if it.status == ST_ERROR)

        if error == 0:
            self._lbl_status.configure(
                text=f"✅  전체 완료 — {done}개 파일 추출 성공", fg=COLOR_SUCCESS)
        else:
            self._lbl_status.configure(
                text=f"⚠  완료 — 성공 {done}개 / 오류 {error}개", fg=COLOR_WARN)

        self._progressbar["value"] = 100
        self._lbl_progress_text.configure(text="")
        self._btn_run.configure(state="normal", text="  텍스트 추출 시작  ")
        self._btn_add.configure(state="normal")
        self._btn_remove.configure(state="normal")
        self._btn_clear.configure(state="normal")
        if self._last_output_folder:
            self._btn_open_folder.configure(state="normal")

    def _on_open_folder(self):
        if self._last_output_folder and os.path.isdir(self._last_output_folder):
            subprocess.Popen(["explorer", self._last_output_folder])

    # ── 헬퍼 ─────────────────────────────────────────────────

    def _update_count(self):
        n = len(self._file_list.items)
        self._lbl_count.configure(text=f"{n}개")

    def _set_run_button_state(self):
        has_items = bool(self._file_list.items)
        self._btn_run.configure(state="normal" if has_items else "disabled")

    def _accent_button(self, parent, text, command, big=False, **kw):
        btn = tk.Button(
            parent, text=text,
            bg=COLOR_ACCENT, fg="white",
            activebackground=COLOR_ACCENT_H, activeforeground="white",
            font=("Segoe UI", 10 if not big else 11, "bold"),
            relief="flat", bd=0,
            padx=12, pady=6 if not big else 9,
            cursor="hand2", command=command, **kw,
        )
        btn.bind("<Enter>", lambda e: btn.configure(
            bg=COLOR_ACCENT_H if btn["state"] != "disabled" else COLOR_BORDER))
        btn.bind("<Leave>", lambda e: btn.configure(
            bg=COLOR_ACCENT if btn["state"] != "disabled" else COLOR_BORDER))
        return btn

    def _ghost_button(self, parent, text, command, **kw):
        btn = tk.Button(
            parent, text=text,
            bg=COLOR_CARD, fg=COLOR_SUBTEXT,
            activebackground=COLOR_BORDER, activeforeground=COLOR_TEXT,
            font=("Segoe UI", 9),
            relief="flat", bd=0,
            padx=10, pady=6,
            cursor="hand2", command=command, **kw,
        )
        btn.bind("<Enter>", lambda e: btn.configure(bg=COLOR_BORDER, fg=COLOR_TEXT))
        btn.bind("<Leave>", lambda e: btn.configure(bg=COLOR_CARD,   fg=COLOR_SUBTEXT))
        return btn


# ── 진입점 ────────────────────────────────────────────────────

if __name__ == "__main__":
    if sys.platform == "win32":
        import ctypes
        try:
            ctypes.windll.kernel32.FreeConsole()
        except Exception:
            pass

    app = PPTXExtractorApp()
    app.mainloop()
